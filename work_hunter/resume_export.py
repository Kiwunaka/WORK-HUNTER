"""Single-column CV export and deterministic text round-trip verification."""
from __future__ import annotations

import hashlib
import html
import re
from pathlib import Path
from typing import Any

from .models import Resume


def resume_lines(body: str) -> list[tuple[int, str]]:
    lines = []
    for source in body.splitlines():
        text = source.strip()
        if not text:
            continue
        heading = re.match(r"^(#{1,3})\s+", text)
        level = len(heading[1]) if heading else 0
        text = text[heading.end():] if heading else text
        # Preserve link destinations as readable text for applicant systems.
        text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
        text = text.replace("**", "").replace("__", "").replace("`", "")
        lines.append((level, text))
    return lines


def extract_resume_text(path: Path) -> str:
    if path.stat().st_size > 10 * 1024 * 1024:
        raise ValueError("Resume file exceeds 10 MiB")
    if path.suffix.lower() == ".docx":
        from docx import Document

        document = Document(str(path))
        return "\n".join(block.text if hasattr(block, "text") else "\n".join(
            " | ".join(cell.text for cell in row.cells) for row in block.rows
        ) for block in document.iter_inner_content())
    if path.suffix.lower() == ".pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if path.suffix.lower() in {".md", ".txt"}:
        return path.read_text(encoding="utf-8-sig")
    raise ValueError("Resume must be PDF, DOCX or text")


def validate_export(path: Path, lines: list[tuple[int, str]]) -> dict[str, Any]:
    extracted = extract_resume_text(path)
    normalized = re.sub(r"\s+", "", extracted)
    cursor = 0
    missing = []
    for _, text in lines:
        needle = re.sub(r"\s+", "", text)
        found = normalized.find(needle, cursor)
        if found < 0:
            missing.append(text)
        else:
            cursor = found + len(needle)
    return {"status": "ok" if not missing else "error", "text": extracted,
            "missing_or_reordered": missing, "scope": "text_round_trip"}


def export_resume(root: Path, resume: Resume, format: str) -> dict[str, Any]:
    if format not in {"pdf", "docx"}:
        raise ValueError("Export format must be pdf or docx")
    lines = resume_lines(resume.body)
    if not lines:
        raise ValueError("Сначала добавьте текст резюме")
    digest = hashlib.sha256(resume.body.encode("utf-8")).hexdigest()
    directory = root / ".work-hunter" / "artifacts" / "exports"
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"resume-{resume.id}-{digest[:16]}.{format}"
    if format == "docx":
        from docx import Document
        from docx.shared import Cm, Pt, RGBColor

        document = Document()
        section = document.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(1.8)
        section.left_margin = section.right_margin = Cm(2)
        for name, size in (("Normal", 10.5), ("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
            style = document.styles[name]
            style.font.name = "Arial"
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.space_after = Pt(6)
        for level, text in lines:
            paragraph = document.add_paragraph(text, style={1: "Title", 2: "Heading 1", 3: "Heading 2"}.get(level, "Normal"))
            paragraph.paragraph_format.keep_together = True
        document.save(str(path))
    else:
        # The app already uses Chromium; printing escaped text needs no network or scripts.
        from playwright.sync_api import sync_playwright

        content = "".join(f"<{tag}>{html.escape(text)}</{tag}>" for level, text in lines
                          for tag in [{1: "h1", 2: "h2", 3: "h3"}.get(level, "p")])
        markup = """<!doctype html><html lang="ru"><meta charset="utf-8"><style>
        @page {size:A4;margin:18mm 20mm} body {font:10.5pt/1.4 Arial,sans-serif;color:#000}
        h1 {font-size:22pt} h2 {font-size:15pt} h3 {font-size:12pt}
        h1,h2,h3 {break-after:avoid;margin:14pt 0 6pt} p {margin:0 0 6pt;white-space:pre-wrap}
        p {orphans:2;widows:2} * {overflow-wrap:anywhere}
        </style><body>""" + content + "</body></html>"
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            try:
                page = browser.new_page(java_script_enabled=False)
                page.route("**/*", lambda route: route.abort())
                page.set_content(markup)
                page.pdf(path=str(path), format="A4", print_background=True, prefer_css_page_size=True)
            finally:
                browser.close()
    validation = validate_export(path, lines)
    return {"path": str(path.resolve()), "format": format,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "validation": validation}
